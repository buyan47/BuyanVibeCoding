"""
Excel / CSV Writer
───────────────────
Takes a list of parsed receipt dicts and writes them to:
  - An Excel workbook (.xlsx) with one sheet per vendor + a combined Summary sheet
  - A flat CSV file for quick inspection

Usage:
    from excel_writer import write_receipts
    write_receipts(records, output_dir="output")
"""

import os
import csv
from typing import List, Dict
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Columns that appear in every sheet (in order)
COMMON_COLUMNS = [
    ("Vendor",        "vendor"),
    ("Date",          "date"),
    ("Amount ($)",    "amount"),
    ("Description",   "description"),
]

# Extra columns per vendor (shown after common columns on vendor-specific sheets)
VENDOR_EXTRA_COLUMNS = {
    "Amtrak": [
        ("Reservation #", "reservation"),
        ("Ticket #",      "ticket_number"),
        ("Passenger",     "passenger"),
        ("Purchase Date", "purchase_date"),
    ],
    "Marriott": [
        ("Check-Out",     "checkout_date"),
        ("Nights",        "nights"),
        ("Confirm #",     "confirmation"),
    ],
    "Uber": [
        ("Trip Fare",     "trip_fare"),
        ("Tip",           "tip"),
        ("Passenger",     "passenger"),
    ],
}

# Header fill colours (hex, no #)
VENDOR_COLOURS = {
    "Amtrak":   "215FA8",    # Amtrak blue
    "Marriott": "B11116",    # Marriott red
    "Uber":     "000000",    # Uber black
    "Summary":  "2E4057",    # neutral dark
}

FONT_COLOUR = "FFFFFF"      # white text on coloured headers


def write_receipts(records: List[Dict], output_dir: str = "output") -> Dict[str, str]:
    """
    Write *records* to Excel and CSV.
    Returns {'excel': path, 'csv': path}.
    """
    os.makedirs(output_dir, exist_ok=True)

    timestamp   = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_path  = os.path.join(output_dir, f"travel_receipts_{timestamp}.xlsx")
    csv_path    = os.path.join(output_dir, f"travel_receipts_{timestamp}.csv")

    _write_excel(records, excel_path)
    _write_csv(records, csv_path)

    print(f"\n[Output] Excel → {excel_path}")
    print(f"[Output] CSV   → {csv_path}")
    return {"excel": excel_path, "csv": csv_path}


# ── Excel ─────────────────────────────────────────────────────────────────────

def _write_excel(records: List[Dict], path: str):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)     # remove default blank sheet

    # Group records by vendor
    vendors = sorted({r.get("vendor", "Unknown") for r in records})
    vendor_groups = {v: [r for r in records if r.get("vendor") == v] for v in vendors}

    # One sheet per vendor
    for vendor, rows in vendor_groups.items():
        columns = COMMON_COLUMNS + VENDOR_EXTRA_COLUMNS.get(vendor, [])
        _add_sheet(wb, vendor, rows, columns, VENDOR_COLOURS.get(vendor, "444444"))

    # Summary sheet (all vendors combined, common columns only)
    if records:
        _add_sheet(wb, "Summary", records, COMMON_COLUMNS, VENDOR_COLOURS["Summary"])
        # Move Summary to first position
        wb.move_sheet("Summary", offset=-len(vendors))

    wb.save(path)


def _add_sheet(wb, sheet_name: str, rows: List[Dict], columns: List, colour: str):
    ws = wb.create_sheet(title=sheet_name)

    header_font    = Font(bold=True, color=FONT_COLOUR)
    header_fill    = PatternFill("solid", fgColor=colour)
    header_align   = Alignment(horizontal="center", vertical="center")
    thin_border    = Border(
        bottom=Side(style="thin", color="CCCCCC"),
        right =Side(style="thin", color="CCCCCC"),
    )

    # Write header row
    for col_idx, (header, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_align

    ws.row_dimensions[1].height = 20

    # Write data rows
    for row_idx, record in enumerate(rows, start=2):
        for col_idx, (_, key) in enumerate(columns, start=1):
            value = record.get(key)
            cell  = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if key == "amount" and value is not None:
                cell.number_format = '"$"#,##0.00'
            cell.alignment = Alignment(vertical="center")

    # Auto-fit column widths
    for col_idx, (header, key) in enumerate(columns, start=1):
        max_len = len(header)
        for record in rows:
            val = record.get(key)
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)

    # Totals row for Amount column
    amount_col = next(
        (i + 1 for i, (_, k) in enumerate(columns) if k == "amount"), None
    )
    if amount_col and rows:
        total_row = len(rows) + 2
        total_cell = ws.cell(row=total_row, column=amount_col)
        total_cell.value  = f"=SUM({get_column_letter(amount_col)}2:{get_column_letter(amount_col)}{total_row-1})"
        total_cell.font   = Font(bold=True)
        total_cell.number_format = '"$"#,##0.00'

        label_cell = ws.cell(row=total_row, column=1, value="TOTAL")
        label_cell.font = Font(bold=True)

    # Freeze header row
    ws.freeze_panes = "A2"


# ── CSV ───────────────────────────────────────────────────────────────────────

def _write_csv(records: List[Dict], path: str):
    if not records:
        return

    # Use union of all keys, with common columns first
    priority = [k for _, k in COMMON_COLUMNS]
    all_keys = priority + [k for r in records for k in r if k not in priority]
    fieldnames = list(dict.fromkeys(all_keys))   # preserve order, deduplicate

    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)
